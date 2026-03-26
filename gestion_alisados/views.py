from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.db.models import Q
from .models import GestionAlisado
from .forms import GestionAlisadoForm
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4, landscape
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from io import BytesIO


def es_staff(user):
    return user.is_staff


@login_required
def lista_gestion_alisados(request):
    """Lista todas las gestiones de alisados registradas con filtros de búsqueda"""
    gestiones = GestionAlisado.objects.all()
    
    # Filtro de búsqueda general
    buscar = request.GET.get('buscar', '')
    if buscar:
        gestiones = gestiones.filter(
            Q(cliente__nombre__icontains=buscar) |
            Q(cliente__apellido__icontains=buscar) |
            Q(procedimiento_realizado_por__icontains=buscar) |
            Q(tipo_alisado__icontains=buscar)
        )
    
    # Filtro por forma natural del cabello
    forma_natural = request.GET.get('forma_natural', '')
    if forma_natural:
        gestiones = gestiones.filter(forma_natural=forma_natural)
    
    # Filtro por porosidad
    porosidad = request.GET.get('porosidad', '')
    if porosidad:
        gestiones = gestiones.filter(porosidad=porosidad)
    
    # Filtro por textura
    textura = request.GET.get('textura', '')
    if textura:
        gestiones = gestiones.filter(textura=textura)
    
    # Filtro por estado de pago
    estado_pago = request.GET.get('estado_pago', '')
    if estado_pago == 'pagado':
        gestiones = gestiones.filter(saldo_pendiente=0)
    elif estado_pago == 'pendiente':
        gestiones = gestiones.filter(saldo_pendiente__gt=0)

    gestiones_firmadas = gestiones.filter(
        firma_consentimiento__isnull=False
    ).exclude(
        firma_consentimiento=''
    )
    
    context = {
        'gestiones': gestiones,
        'gestiones_firmadas': gestiones_firmadas,
    }
    return render(request, 'gestion_alisados/lista_gestion_alisados.html', context)


@login_required
def crear_gestion_alisado(request):
    """Crea un nuevo registro de gestión de alisado"""
    is_modal = request.GET.get('modal') == '1'
    
    if request.method == 'POST':
        print(f"=== CREAR GESTION ALISADO - POST recibido ===")
        print(f"is_modal: {is_modal}")
        print(f"X-Requested-With: {request.headers.get('X-Requested-With')}")
        print(f"POST data keys: {list(request.POST.keys())}")
        print(f"FILES: {list(request.FILES.keys())}")
        
        form = GestionAlisadoForm(request.POST, request.FILES)
        if form.is_valid():
            print("=== FORMULARIO VÁLIDO ===")
            gestion = form.save()
            print(f"=== GESTIÓN GUARDADA con ID: {gestion.pk} ===")
            
            if is_modal or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                # Retornar respuesta JSON para AJAX
                return JsonResponse({
                    'success': True,
                    'message': 'Gestión de alisado registrada exitosamente.'
                })
            
            messages.success(request, 'Gestión de alisado registrada exitosamente.')
            return redirect('gestion_alisados:lista_gestion_alisados')
        else:
            print("=== FORMULARIO INVÁLIDO ===")
            print(f"Errores: {form.errors}")
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                # Si hay errores y es AJAX, devolver JSON con los errores
                return JsonResponse({
                    'success': False,
                    'message': 'Por favor corrija los errores en el formulario.',
                    'errors': form.errors
                }, status=400)
    else:
        form = GestionAlisadoForm()
    
    context = {
        'form': form,
        'titulo': 'Gestión de Alisado',
        'is_modal': is_modal
    }
    
    # Si es modal, usar template simplificado
    if is_modal:
        return render(request, 'gestion_alisados/form_gestion_alisado_modal_content.html', context)
    
    return render(request, 'gestion_alisados/form_gestion_alisado.html', context)


@login_required
@user_passes_test(es_staff)
def ver_gestion_alisado(request, pk):
    """Muestra los detalles de una gestión de alisado"""
    gestion = get_object_or_404(GestionAlisado, pk=pk)
    is_modal = request.GET.get('modal') == '1'
    context = {
        'gestion': gestion,
        'is_modal': is_modal,
    }

    if is_modal:
        return render(request, 'gestion_alisados/detalle_gestion_alisado_modal_content.html', context)

    return render(request, 'gestion_alisados/detalle_gestion_alisado.html', context)


@login_required
def editar_gestion_alisado(request, pk):
    """Edita una gestión de alisado existente"""
    gestion = get_object_or_404(GestionAlisado, pk=pk)
    is_modal = request.GET.get('modal') == '1'
    es_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    
    if request.method == 'POST':
        form = GestionAlisadoForm(request.POST, request.FILES, instance=gestion)
        if form.is_valid():
            gestion = form.save()

            if es_ajax:
                return JsonResponse({
                    'success': True,
                    'message': 'Gestión de alisado actualizada exitosamente.'
                })

            messages.success(request, 'Gestión de alisado actualizada exitosamente.')
            return redirect('gestion_alisados:ver_gestion_alisado', pk=gestion.pk)
        
        if es_ajax:
            return JsonResponse({
                'success': False,
                'message': 'Por favor corrija los errores en el formulario.',
                'errors': form.errors
            }, status=400)
    else:
        form = GestionAlisadoForm(instance=gestion)
    
    context = {
        'form': form,
        'titulo': 'Editar Gestión de Alisado',
        'gestion': gestion,
        'cliente_actual': gestion.cliente,
        'is_modal': is_modal,
    }

    if is_modal:
        return render(request, 'gestion_alisados/form_gestion_alisado_modal_content.html', context)

    return render(request, 'gestion_alisados/form_gestion_alisado.html', context)


@login_required
def eliminar_gestion_alisado(request, pk):
    """Elimina una gestión de alisado"""
    gestion = get_object_or_404(GestionAlisado, pk=pk)
    is_modal = request.GET.get('modal') == '1'
    es_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    
    if request.method == 'POST':
        gestion.delete()

        if es_ajax:
            return JsonResponse({
                'success': True,
                'message': 'Gestión de alisado eliminada exitosamente.'
            })

        messages.success(request, 'Gestión de alisado eliminada exitosamente.')
        return redirect('gestion_alisados:lista_gestion_alisados')
    
    context = {
        'gestion': gestion,
        'is_modal': is_modal,
    }

    if is_modal:
        return render(request, 'gestion_alisados/eliminar_gestion_alisado_modal_content.html', context)

    return render(request, 'gestion_alisados/eliminar_gestion_alisado.html', context)


@login_required
def exportar_excel(request):
    """Exporta todas las gestiones de alisados a un archivo Excel"""
    # Obtener las gestiones con los mismos filtros que en lista_gestion_alisados
    gestiones = GestionAlisado.objects.all()
    
    # Aplicar filtros si existen
    buscar = request.GET.get('buscar', '')
    if buscar:
        gestiones = gestiones.filter(
            Q(cliente__nombre__icontains=buscar) |
            Q(cliente__apellido__icontains=buscar) |
            Q(procedimiento_realizado_por__icontains=buscar) |
            Q(tipo_alisado__icontains=buscar)
        )
    
    forma_natural = request.GET.get('forma_natural', '')
    if forma_natural:
        gestiones = gestiones.filter(forma_natural=forma_natural)
    
    porosidad = request.GET.get('porosidad', '')
    if porosidad:
        gestiones = gestiones.filter(porosidad=porosidad)
    
    textura = request.GET.get('textura', '')
    if textura:
        gestiones = gestiones.filter(textura=textura)
    
    estado_pago = request.GET.get('estado_pago', '')
    if estado_pago == 'pagado':
        gestiones = gestiones.filter(saldo_pendiente=0)
    elif estado_pago == 'pendiente':
        gestiones = gestiones.filter(saldo_pendiente__gt=0)
    
    # Crear un nuevo libro de trabajo
    wb = Workbook()
    ws = wb.active
    ws.title = "Gestión de Alisados"
    
    # Definir estilos
    header_fill = PatternFill(start_color="3a2a24", end_color="3a2a24", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # Headers
    headers = [
        'Fecha y Hora', 'Cliente', 'Documento', 'Profesional', 'Tipo de Alisado',
        'Precio', 'Anticipo', 'Saldo Pendiente', 'Porosidad', 'Textura', 'Forma Natural'
    ]
    ws.append(headers)
    
    # Aplicar estilos al header
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = border
    
    # Agregar datos
    for gestion in gestiones:
        cliente_nombre = f"{gestion.cliente.nombre} {gestion.cliente.apellido}"
        ws.append([
            gestion.fecha_hora.strftime('%d/%m/%Y %H:%M') if gestion.fecha_hora else '',
            cliente_nombre,
            gestion.cliente.numero_documento or '',
            gestion.procedimiento_realizado_por or '',
            gestion.tipo_alisado or '',
            f"${float(gestion.precio_alisado):.2f}" if gestion.precio_alisado else '',
            f"${float(gestion.anticipo_cliente):.2f}" if gestion.anticipo_cliente else '',
            f"${float(gestion.saldo_pendiente):.2f}" if gestion.saldo_pendiente else '',
            gestion.get_porosidad_display() or '',
            gestion.get_textura_display() or '',
            gestion.get_forma_natural_display() or ''
        ])
    
    # Ajustar ancho de columnas
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 15
    ws.column_dimensions['I'].width = 12
    ws.column_dimensions['J'].width = 12
    ws.column_dimensions['K'].width = 15
    
    # Aplicar bordes a las celdas de datos
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=11):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='center')
    
    # Crear la respuesta
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="gestion_alisados_{datetime.now().strftime("%d%m%Y_%H%M%S")}.xlsx"'
    wb.save(response)
    
    return response


@login_required
def exportar_pdf(request):
    """Exporta todas las gestiones de alisados a un archivo PDF en horizontal"""
    # Obtener las gestiones con los mismos filtros que en lista_gestion_alisados
    gestiones = GestionAlisado.objects.all()
    
    # Aplicar filtros si existen
    buscar = request.GET.get('buscar', '')
    if buscar:
        gestiones = gestiones.filter(
            Q(cliente__nombre__icontains=buscar) |
            Q(cliente__apellido__icontains=buscar) |
            Q(procedimiento_realizado_por__icontains=buscar) |
            Q(tipo_alisado__icontains=buscar)
        )
    
    forma_natural = request.GET.get('forma_natural', '')
    if forma_natural:
        gestiones = gestiones.filter(forma_natural=forma_natural)
    
    porosidad = request.GET.get('porosidad', '')
    if porosidad:
        gestiones = gestiones.filter(porosidad=porosidad)
    
    textura = request.GET.get('textura', '')
    if textura:
        gestiones = gestiones.filter(textura=textura)
    
    estado_pago = request.GET.get('estado_pago', '')
    if estado_pago == 'pagado':
        gestiones = gestiones.filter(saldo_pendiente=0)
    elif estado_pago == 'pendiente':
        gestiones = gestiones.filter(saldo_pendiente__gt=0)
    
    # Crear el PDF en memoria HORIZONTAL
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=0.5*inch, bottomMargin=0.5*inch)
    story = []
    
    # Estilos
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#3a2a24'),
        spaceAfter=20,
        alignment=1  # Centrado
    )
    
    # Título
    story.append(Paragraph('Gestión de Alisados', title_style))
    story.append(Spacer(1, 0.3*inch))
    
    # Preparar datos de la tabla
    table_data = [[
        'Fecha y Hora', 'Cliente', 'Documento', 'Profesional', 'Tipo de Alisado',
        'Precio', 'Anticipo', 'Saldo', 'Porosidad', 'Textura'
    ]]
    
    for gestion in gestiones:
        cliente_nombre = f"{gestion.cliente.nombre} {gestion.cliente.apellido}"
        table_data.append([
            gestion.fecha_hora.strftime('%d/%m/%Y %H:%M') if gestion.fecha_hora else '',
            cliente_nombre,
            gestion.cliente.numero_documento or '',
            gestion.procedimiento_realizado_por or '',
            gestion.tipo_alisado or '',
            f"${float(gestion.precio_alisado):.2f}" if gestion.precio_alisado else '',
            f"${float(gestion.anticipo_cliente):.2f}" if gestion.anticipo_cliente else '',
            f"${float(gestion.saldo_pendiente):.2f}" if gestion.saldo_pendiente else '',
            gestion.get_porosidad_display() or '',
            gestion.get_textura_display() or ''
        ])
    
    # Crear tabla con mejor ancho para horizontal
    table = Table(table_data, colWidths=[1.0*inch, 1.2*inch, 1.0*inch, 1.0*inch, 1.2*inch, 
                                         0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch])
    
    # Estilos para la tabla
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3a2a24')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f0eb')]),
    ]))
    
    story.append(table)
    
    # Generar PDF
    doc.build(story)
    buffer.seek(0)
    
    # Crear respuesta
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="gestion_alisados_{datetime.now().strftime("%d%m%Y_%H%M%S")}.pdf"'
    
    return response


@login_required
def imprimir_terminos_firmados(request):
    """Vista imprimible con solo consentimientos firmados."""
    gestiones_base = GestionAlisado.objects.select_related('cliente').all()

    buscar = request.GET.get('buscar', '')
    if buscar:
        gestiones_base = gestiones_base.filter(
            Q(cliente__nombre__icontains=buscar) |
            Q(cliente__apellido__icontains=buscar) |
            Q(procedimiento_realizado_por__icontains=buscar) |
            Q(tipo_alisado__icontains=buscar)
        )

    forma_natural = request.GET.get('forma_natural', '')
    if forma_natural:
        gestiones_base = gestiones_base.filter(forma_natural=forma_natural)

    porosidad = request.GET.get('porosidad', '')
    if porosidad:
        gestiones_base = gestiones_base.filter(porosidad=porosidad)

    textura = request.GET.get('textura', '')
    if textura:
        gestiones_base = gestiones_base.filter(textura=textura)

    estado_pago = request.GET.get('estado_pago', '')
    if estado_pago == 'pagado':
        gestiones_base = gestiones_base.filter(saldo_pendiente=0)
    elif estado_pago == 'pendiente':
        gestiones_base = gestiones_base.filter(saldo_pendiente__gt=0)

    gestiones = gestiones_base.filter(
        firma_consentimiento__isnull=False
    ).exclude(
        firma_consentimiento=''
    ).order_by('-fecha_hora')

    context = {
        'gestiones': gestiones,
        'total_firmadas': gestiones.count(),
        'autoprint': request.GET.get('autoprint') == '1',
        'fecha_impresion': datetime.now(),
    }
    return render(request, 'gestion_alisados/imprimir_terminos_firmados.html', context)


@login_required
def imprimir_consentimiento_individual(request, pk):
    """Vista imprimible de consentimiento para una gestión específica."""
    gestion = get_object_or_404(GestionAlisado.objects.select_related('cliente'), pk=pk)
    context = {
        'gestion': gestion,
        'autoprint': request.GET.get('autoprint') == '1',
        'fecha_impresion': datetime.now(),
        'tiene_firma': bool(gestion.firma_consentimiento),
    }
    return render(request, 'gestion_alisados/imprimir_consentimiento_individual.html', context)
