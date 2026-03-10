from io import BytesIO

from django.http import HttpResponse
from django.utils.text import slugify
from django.contrib.staticfiles import finders

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    Image as RLImage,
)

PAGE_WIDTH, PAGE_HEIGHT = A4


def _nombre_archivo_compra(compra, extension):
    proveedor = getattr(compra.proveedor, "nombre_proveedor", "proveedor")
    proveedor_slug = slugify(proveedor) or "proveedor"
    return f"comprobante_compra_{compra.id}_{proveedor_slug}.{extension}"


def _fmt_money(value):
    try:
        return f"${int(value):,}".replace(",", ".")
    except Exception:
        return "$0"


def _find_first_static(paths):
    for path in paths:
        found = finders.find(path)
        if found:
            return found
    return None


def _get_logo_path():
    return _find_first_static([
        "compras/img/logo_monakeratina.png",
        "compras/img/logo_monakeratina.webp",
        "compras/img/logo_monakeratina.jpg",
        "compras/img/logo_monakefratina.png",
        "compras/img/logo_monakefratina.webp",
        "compras/img/logo_monakefratina.jpg",
        "compras/img/logo_monakeratina_watermark.png",
        "compras/img/logo_monakefratina_watermark.png",
        "core/img/logo_monakeratina.png",
        "core/img/logo_monakefratina.png",
        "core/img/logo_monakeratina_watermark.png",
        "core/img/logo_monakefratina_watermark.png",
    ])


def _get_watermark_path():
    return _find_first_static([
        "compras/img/logo_monakeratina_watermark.png",
        "compras/img/logo_monakefratina_watermark.png",
        "compras/img/logo_monakeratina.png",
        "compras/img/logo_monakefratina.png",
        "core/img/logo_monakeratina_watermark.png",
        "core/img/logo_monakefratina_watermark.png",
        "core/img/logo_monakeratina.png",
        "core/img/logo_monakefratina.png",
    ])


def _draw_watermark(canvas, doc):
    watermark_path = _get_watermark_path()
    if not watermark_path:
        return

    canvas.saveState()

    try:
        if hasattr(canvas, "setFillAlpha"):
            canvas.setFillAlpha(0.06)
    except Exception:
        pass

    image_width = 18 * mm
    image_height = 18 * mm
    gap_x = 16 * mm
    gap_y = 18 * mm

    x = 10 * mm
    while x < PAGE_WIDTH:
        y = 12 * mm
        while y < PAGE_HEIGHT:
            canvas.drawImage(
                watermark_path,
                x,
                y,
                width=image_width,
                height=image_height,
                preserveAspectRatio=True,
                mask="auto",
            )
            y += image_height + gap_y
        x += image_width + gap_x

    canvas.restoreState()


def build_comprobante_pdf_response(compra):
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#000000"),
        alignment=1,  # centrado
        spaceAfter=2,
    )

    small_label_style = ParagraphStyle(
        "SmallLabelStyle",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=9,
        textColor=colors.HexColor("#333333"),
    )

    small_value_style = ParagraphStyle(
        "SmallValueStyle",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=13,
        textColor=colors.HexColor("#111111"),
    )

    text_style = ParagraphStyle(
        "TextStyle",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#222222"),
    )

    footer_style = ParagraphStyle(
        "FooterStyle",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#666666"),
        alignment=1,
    )

    elements = []

    proveedor = getattr(compra.proveedor, "nombre_proveedor", "—")
    usuario = compra.usuario.username if compra.usuario else "—"
    fecha = compra.fecha.strftime("%d/%m/%Y") if compra.fecha else "—"
    estado = "Anulada" if compra.anulada else "Activa"
    total = compra.precio_total or 0

    # Header con logo + título
    logo_path = _get_logo_path()
    logo_flowable = ""
    if logo_path:
        try:
            logo_flowable = RLImage(logo_path, width=34 * mm, height=16 * mm)
        except Exception:
            logo_flowable = ""

    header_table = Table(
        [[logo_flowable, Paragraph("Comprobante", title_style), ""]],
        colWidths=[42 * mm, 100 * mm, 26 * mm],
    )
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, 0), "LEFT"),
        ("ALIGN", (1, 0), (1, 0), "CENTER"),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.black),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 8))

    # Primera línea tipo factura
    numero_box = Table([
        [Paragraph("NÚMERO", small_label_style)],
        [Paragraph(f"{compra.id:04d}", small_value_style)],
    ], colWidths=[28 * mm])

    de_box = Table([
        [Paragraph("DE", small_label_style)],
        [Paragraph("Monakeratina", small_value_style)],
        [Paragraph("Sistema de compras", text_style)],
        [Paragraph("Documento interno", text_style)],
    ], colWidths=[64 * mm])

    para_box = Table([
        [Paragraph("PARA", small_label_style)],
        [Paragraph(str(proveedor), small_value_style)],
        [Paragraph("Proveedor asociado a la compra", text_style)],
    ], colWidths=[74 * mm])

    top_info = Table(
        [[numero_box, de_box, para_box]],
        colWidths=[30 * mm, 66 * mm, 76 * mm],
    )
    top_info.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEABOVE", (0, 0), (-1, -1), 1.0, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(top_info)
    elements.append(Spacer(1, 10))

    # Segunda línea meta
    fecha_box = Table([
        [Paragraph("FECHA", small_label_style)],
        [Paragraph(fecha, small_value_style)],
    ], colWidths=[28 * mm])

    usuario_box = Table([
        [Paragraph("USUARIO", small_label_style)],
        [Paragraph(usuario, small_value_style)],
    ], colWidths=[28 * mm])

    estado_box = Table([
        [Paragraph("ESTADO", small_label_style)],
        [Paragraph(estado, small_value_style)],
    ], colWidths=[28 * mm])

    meta_table = Table(
        [[fecha_box, usuario_box, estado_box]],
        colWidths=[30 * mm, 30 * mm, 30 * mm],
    )
    meta_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEABOVE", (0, 0), (-1, -1), 1.0, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 10))

    # Tabla de productos
    table_data = [["Descripción", "Cantidad", "Precio unidad", "Importe"]]

    detalles = list(compra.detalles.all())
    if detalles:
        for detalle in detalles:
            subtotal = getattr(detalle, "subtotal", 0) or 0
            table_data.append([
                str(detalle.producto.nombre),
                str(detalle.cantidad),
                _fmt_money(detalle.precio_unitario),
                _fmt_money(subtotal),
            ])
    else:
        table_data.append(["No hay productos en esta compra.", "", "", ""])

    detail_table = Table(
        table_data,
        colWidths=[88 * mm, 26 * mm, 32 * mm, 32 * mm],
        repeatRows=1,
    )
    detail_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 1.0, colors.black),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#777777")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("ALIGN", (1, 1), (1, -1), "CENTER"),
        ("ALIGN", (2, 1), (3, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 1), (-1, -1), 0.35, colors.HexColor("#CFCFCF")),
    ]))
    elements.append(detail_table)
    elements.append(Spacer(1, 6))

    # Total abajo a la derecha
    total_table = Table(
        [["Total", _fmt_money(total)]],
        colWidths=[34 * mm, 34 * mm],
        hAlign="RIGHT",
    )
    total_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 1.0, colors.black),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (0, 0), (0, 0), "LEFT"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(total_table)

    elements.append(Spacer(1, 10))
    elements.append(Paragraph("Generado por Monakeratina", footer_style))

    doc.build(
        elements,
        onFirstPage=_draw_watermark,
        onLaterPages=_draw_watermark,
    )

    pdf = buffer.getvalue()
    buffer.close()

    filename = _nombre_archivo_compra(compra, "pdf")
    response = HttpResponse(pdf, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def build_comprobante_excel_response(compra):
    wb = Workbook()
    ws = wb.active
    ws.title = "Comprobante"

    title_font = Font(bold=True, size=14, color="111111")
    header_font = Font(bold=True, color="FFFFFF")
    label_font = Font(bold=True, color="111111")
    total_font = Font(bold=True, color="111111")

    header_fill = PatternFill("solid", fgColor="111111")
    soft_fill = PatternFill("solid", fgColor="F3F4F6")
    light_fill = PatternFill("solid", fgColor="F8FAFC")

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    proveedor = getattr(compra.proveedor, "nombre_proveedor", "—")
    usuario = compra.usuario.username if compra.usuario else "—"
    fecha = compra.fecha.strftime("%d/%m/%Y") if compra.fecha else "—"
    estado = "Anulada" if compra.anulada else "Activa"
    total = compra.precio_total or 0

    logo_path = _get_logo_path()
    if logo_path:
        try:
            img = XLImage(logo_path)
            img.width = 180
            img.height = 55
            ws.add_image(img, "A1")
        except Exception:
            pass

    ws.merge_cells("A1:D2")
    ws["A1"] = f"Comprobante de Compra #{compra.id}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    info_rows = [
        ("Proveedor", proveedor),
        ("Usuario", usuario),
        ("Fecha", fecha),
        ("Estado", estado),
        ("Total", float(total or 0)),
    ]

    start_info_row = 4
    for idx, (label, value) in enumerate(info_rows, start=start_info_row):
        label_cell = ws[f"A{idx}"]
        value_cell = ws[f"B{idx}"]

        label_cell.value = label
        label_cell.font = label_font
        label_cell.fill = light_fill
        label_cell.border = thin_border
        label_cell.alignment = Alignment(vertical="center")

        value_cell.value = value
        value_cell.border = thin_border
        value_cell.alignment = Alignment(vertical="center")

    ws[f"B{start_info_row + 4}"].number_format = '$#,##0'

    table_start = start_info_row + len(info_rows) + 2

    headers = ["Producto", "Cantidad", "Precio Unitario", "Subtotal"]
    for col_num, header in enumerate(headers, start=1):
        cell = ws.cell(row=table_start, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    current_row = table_start + 1
    detalles = list(compra.detalles.all())

    if detalles:
        for detalle in detalles:
            subtotal = getattr(detalle, "subtotal", 0) or 0

            ws.cell(row=current_row, column=1, value=str(detalle.producto.nombre))
            ws.cell(row=current_row, column=2, value=detalle.cantidad)
            ws.cell(row=current_row, column=3, value=float(detalle.precio_unitario or 0))
            ws.cell(row=current_row, column=4, value=float(subtotal or 0))

            for col in range(1, 5):
                cell = ws.cell(row=current_row, column=col)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

            ws.cell(row=current_row, column=3).number_format = '$#,##0'
            ws.cell(row=current_row, column=4).number_format = '$#,##0'
            current_row += 1
    else:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        empty_cell = ws.cell(row=current_row, column=1, value="No hay productos en esta compra.")
        empty_cell.border = thin_border
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += 1

    ws.cell(row=current_row, column=3, value="TOTAL").font = total_font
    ws.cell(row=current_row, column=4, value=float(total or 0)).font = total_font
    ws.cell(row=current_row, column=4).number_format = '$#,##0'

    ws.cell(row=current_row, column=3).fill = soft_fill
    ws.cell(row=current_row, column=4).fill = soft_fill
    ws.cell(row=current_row, column=3).border = thin_border
    ws.cell(row=current_row, column=4).border = thin_border
    ws.cell(row=current_row, column=3).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=current_row, column=4).alignment = Alignment(vertical="center")

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = _nombre_archivo_compra(compra, "xlsx")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response