from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Q, Sum, Value, DecimalField
from django.db.models.functions import Coalesce
from .models import Venta, DetalleVenta, DevolucionVenta, DetalleDevolucion
from inventario.models import Stock
from django.db.models import F
from compras.models import DetalleCompra
from django.contrib import messages
from .forms import VentaForm
from Productos.models import Producto
from django.core.exceptions import ValidationError
import json
from django.db import transaction
from servicios.models import Servicio
from django.http import JsonResponse
from decimal import Decimal
from django.views.decorators.http import require_POST
from django.template.loader import render_to_string
from django.apps import apps
from django.db.models.functions import TruncDate
from datetime import datetime
from io import BytesIO
from django.http import HttpResponse
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, LineChart, Reference
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer, Paragraph, Image as RLImage
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

Personal = apps.get_model("personal", "Personal")


def lista_ventas(request):
    q = request.GET.get("q", "").strip()
    estado = request.GET.get("estado", "activa").strip()
    sort = request.GET.get("sort", "fecha").strip()
    direction = request.GET.get("dir", "desc").strip()

    ventas = (
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__producto", "detalles__servicio")
        .annotate(
            total_orden=Coalesce(
                Sum("detalles__subtotal"),
                Value(0),
                output_field=DecimalField(max_digits=12, decimal_places=2),
            )
        )
        .all()
    )

    if q:
        ventas = ventas.filter(
            Q(codigo_venta__icontains=q)
            | Q(cliente__nombre__icontains=q)
            | Q(cliente__apellido__icontains=q)
        )

    if estado:
        ventas = ventas.filter(estado=estado)

    # ==================== ORDENAMIENTO ====================
    if direction not in ["asc", "desc"]:
        direction = "desc"

    if sort == "codigo_venta":
        ventas = ventas.order_by(
            "codigo_venta" if direction == "asc" else "-codigo_venta"
        )

    elif sort == "cliente":
        ventas = ventas.order_by(
            "cliente__nombre" if direction == "asc" else "-cliente__nombre",
            "cliente__apellido" if direction == "asc" else "-cliente__apellido",
        )

    elif sort == "estado":
        ventas = ventas.order_by("estado" if direction == "asc" else "-estado")

    elif sort == "total":
        ventas = ventas.order_by(
            "total_orden" if direction == "asc" else "-total_orden"
        )

    else:
        sort = "fecha"
        ventas = ventas.order_by("fecha" if direction == "asc" else "-fecha")

    return render(
        request,
        "ventas/lista_ventas.html",
        {
            "ventas": ventas,
            "q": q,
            "estado": estado,
            "sort": sort,
            "direction": direction,
        },
    )


@require_POST
def toggle_estado_venta(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.prefetch_related("detalles__producto", "detalles__devoluciones"),
        id=venta_id
    )
    nuevo_estado = "anulada" if venta.estado == "activa" else "activa"

    with transaction.atomic():
        for detalle in venta.detalles.all():
            if not detalle.producto:
                continue
            if nuevo_estado == "anulada":
                # Restaurar solo lo que no fue devuelto aún
                cantidad_restante = detalle.cantidad_disponible_devolver
                if cantidad_restante > 0:
                    Stock.objects.filter(producto=detalle.producto).update(
                        cantidad_actual=F("cantidad_actual") + cantidad_restante
                    )
            else:
                # Vuelve a activa: descontar lo que no había sido devuelto
                cantidad_restante = detalle.cantidad_disponible_devolver
                if cantidad_restante > 0:
                    Stock.objects.filter(producto=detalle.producto).update(
                        cantidad_actual=F("cantidad_actual") - cantidad_restante
                    )
        venta.estado = nuevo_estado
        venta.save()

    messages.success(
        request, f"Estado actualizado correctamente: {venta.estado.upper()}"
    )
    from django.urls import reverse
    return redirect(reverse("ventas:lista") + "?estado=activa")


def es_ajax(request):
    return request.headers.get("x-requested-with") == "XMLHttpRequest"


def render_crear_venta(request, form, productos_stock, servicios, personal, status=200):
    """
    - Si es AJAX: retorna HTML parcial para meterlo dentro del modal.
    - Si no es AJAX: retorna la página completa normal.
    """
    if es_ajax(request):
        html = render_to_string(
            "ventas/partials/crear_venta_form.html",
            {
                "form": form,
                "productos_stock": productos_stock,
                "servicios": servicios,
                "personal": personal,
            },
            request=request,
        )
        return JsonResponse({"success": False, "html": html}, status=status)

    return render(
        request,
        "ventas/crear_venta.html",
        {
            "form": form,
            "productos_stock": productos_stock,
            "servicios": servicios,
            "personal": personal,
        },
        status=status,
    )


@transaction.atomic
def crear_venta(request):
    print(">>> ENTRÓ A crear_venta (views.py correcto)")
    productos = Producto.objects.all()
    productos_stock = []

    for p in productos:
        entradas = (
            DetalleCompra.objects.filter(producto=p).aggregate(total=Sum("cantidad"))[
                "total"
            ]
            or 0
        )

        salidas = (
            DetalleVenta.objects.filter(producto=p, venta__estado="activa").aggregate(
                total=Sum("cantidad")
            )["total"]
            or 0
        )

        devuelto = (
            DetalleDevolucion.objects.filter(
                detalle_venta__producto=p,
                detalle_venta__venta__estado="activa",
            ).aggregate(total=Sum("cantidad_devuelta"))["total"]
            or 0
        )

        stock_real = entradas - salidas + devuelto

        productos_stock.append({"producto": p, "stock": stock_real})

    servicios = Servicio.objects.all()
    personal = Personal.objects.filter(rol="Colaborador", activo=True).order_by(
        "nombres", "apellidos"
    )
    if request.method == "POST":
        print(">>> POST LLEGÓ")
        form = VentaForm(request.POST)
        items_json = request.POST.get("items")

        print(">>> items_json RAW:", items_json)
        print(">>> form.is_valid:", form.is_valid())
        print(">>> form.errors:", form.errors)

        # 1) Validar items_json
        if not items_json:
            messages.error(
                request,
                "No se recibieron items. Agrega al menos 1 producto o servicio.",
            )
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        # 2) Convertir JSON a lista
        try:
            items = json.loads(items_json)
        except json.JSONDecodeError:
            messages.error(request, "El JSON de items llegó dañado. Revisa ventas.js.")
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        print(">>> items PARSEADOS:", items)

        # 3) Debe haber items
        if not items:
            messages.error(
                request, "Agrega al menos 1 producto o servicio antes de guardar."
            )
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        # 4) Validar formulario
        if not form.is_valid():
            messages.error(request, "Formulario inválido. Revisa los campos.")
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        # ===============================
        # CREAR VENTA
        # ===============================
        venta = form.save(commit=False)
        venta.codigo_colaborador = "PENDIENTE"
        venta.nombre_colaborador = "Pendiente"

        # ✅ codigo_producto es obligatorio en tu modelo Venta
        # lo llenamos con un resumen simple
        it0 = items[0]
        if it0.get("tipo") == "producto":
            base = it0.get("id", "")
        else:
            base = it0.get("id_servicio", "")
        venta.codigo_producto = (
            str(base) if len(items) == 1 else f"{base} (+{len(items) - 1})"
        )

        venta.save()

        # ===============================
        # GUARDAR DETALLES
        # ===============================
        for item in items:
            if item.get("tipo") == "producto":
                codigo = item.get("id")
                if not codigo:
                    continue
                    entradas = (
                        DetalleCompra.objects.filter(producto=Producto).aggregate(
                            total=Sum("cantidad")
                        )["total"]
                        or 0
                    )

                producto = Producto.objects.select_for_update().get(codigo=codigo)

                entradas = (
                    DetalleCompra.objects.filter(producto=producto).aggregate(
                        total=Sum("cantidad")
                    )["total"]
                    or 0
                )
                salidas = (
                    DetalleVenta.objects.filter(
                        producto=producto, venta__estado="activa"
                    ).aggregate(total=Sum("cantidad"))["total"]
                    or 0
                )
                devuelto_prod = (
                    DetalleDevolucion.objects.filter(
                        detalle_venta__producto=producto,
                        detalle_venta__venta__estado="activa",
                    ).aggregate(total=Sum("cantidad_devuelta"))["total"]
                    or 0
                )
                stock_real = entradas - salidas + devuelto_prod

                if int(item["cantidad"]) > stock_real:
                    raise ValidationError(f"Stock insuficiente para {producto.nombre}")

                DetalleVenta.objects.create(
                    venta=venta,
                    producto=producto,
                    precio_unitario=Decimal(str(item["precio"])),
                    cantidad=int(item["cantidad"]),
                    subtotal=Decimal(str(item["subtotal"])),
                )
                # Descontar del Stock
                Stock.objects.filter(producto=producto).update(
                    cantidad_actual=F("cantidad_actual") - int(item["cantidad"])
                )

            elif item.get("tipo") == "servicio":
                servicio = Servicio.objects.get(id_servicio=item["id_servicio"])
                colaborador = Personal.objects.get(id=item["id_personal"])

                DetalleVenta.objects.create(
                    venta=venta,
                    servicio=servicio,
                    colaborador_servicio=colaborador,
                    precio_unitario=Decimal(str(item["precio"])),
                    cantidad=int(item.get("cantidad", 1)),
                    subtotal=Decimal(str(item.get("subtotal", item["precio"]))),
                )

        if es_ajax(request):
            return JsonResponse({"success": True})

        messages.success(request, "Venta registrada correctamente.")
        return redirect("ventas:lista")

    # GET
    form = VentaForm()
    if es_ajax(request):
        html = render_to_string(
            "ventas/partials/crear_venta_form.html",
            {
                "form": form,
                "productos_stock": productos_stock,
                "servicios": servicios,
                "personal": personal,
            },
            request=request,
        )
        return JsonResponse({"success": True, "html": html})

    return render(
        request,
        "ventas/crear_venta.html",
        {
            "form": form,
            "productos_stock": productos_stock,
            "servicios": servicios,
            "personal": personal,
        },
    )


def editar_venta_modal(request, pk):
    venta = get_object_or_404(
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__producto", "detalles__devoluciones", "detalles__servicio"),
        pk=pk,
    )

    detalles_productos = venta.detalles.filter(producto__isnull=False).select_related("producto")
    detalles_servicios = venta.detalles.filter(servicio__isnull=False).select_related("servicio")
    servicios = Servicio.objects.all()
    personal = Personal.objects.filter(rol="Colaborador", activo=True).order_by("nombres", "apellidos")

    # Stock disponible para cada producto de la venta (excluyendo esta venta)
    stock_por_producto = {}
    for det in detalles_productos:
        p = det.producto
        entradas = DetalleCompra.objects.filter(producto=p).aggregate(t=Sum("cantidad"))["t"] or 0
        salidas = (
            DetalleVenta.objects.filter(producto=p, venta__estado="activa")
            .exclude(venta=venta)
            .aggregate(t=Sum("cantidad"))["t"] or 0
        )
        devuelto = (
            DetalleDevolucion.objects.filter(
                detalle_venta__producto=p,
                detalle_venta__venta__estado="activa",
            ).aggregate(t=Sum("cantidad_devuelta"))["t"] or 0
        )
        # Stock real + lo que ya tiene esta venta (porque se va a editar)
        ya_tiene = det.cantidad_disponible_devolver
        stock_por_producto[det.id] = entradas - salidas + devuelto + ya_tiene

    if request.method == "POST":
        errores = []
        with transaction.atomic():
            # 1) Restaurar stock de los productos actuales
            for det in detalles_productos:
                Stock.objects.filter(producto=det.producto).update(
                    cantidad_actual=F("cantidad_actual") + det.cantidad_disponible_devolver
                )

            # 2) Guardar cambios de productos
            for det in detalles_productos:
                cant_str = request.POST.get(f"prod_cant_{det.id}")
                precio_str = request.POST.get(f"prod_precio_{det.id}")
                nuevo_codigo = request.POST.get(f"prod_codigo_{det.id}")
                if cant_str is None or precio_str is None:
                    continue

                nueva_cantidad = int(cant_str)
                nuevo_precio = Decimal(precio_str)

                # Cambio de producto si seleccionó uno diferente
                if nuevo_codigo and nuevo_codigo != det.producto.codigo:
                    det.producto = Producto.objects.get(codigo=nuevo_codigo)

                # Validar stock del producto (puede ser el nuevo)
                ent = DetalleCompra.objects.filter(producto=det.producto).aggregate(t=Sum("cantidad"))["t"] or 0
                sal = (
                    DetalleVenta.objects.filter(producto=det.producto, venta__estado="activa")
                    .exclude(venta=venta)
                    .aggregate(t=Sum("cantidad"))["t"] or 0
                )
                dev = (
                    DetalleDevolucion.objects.filter(
                        detalle_venta__producto=det.producto,
                        detalle_venta__venta__estado="activa",
                    ).aggregate(t=Sum("cantidad_devuelta"))["t"] or 0
                )
                disponible = ent - sal + dev
                if nueva_cantidad > disponible:
                    nombre = det.producto.nombre
                    return JsonResponse(
                        {"ok": False, "error": f"Stock insuficiente para '{nombre}'. Disponible: {disponible}."},
                        status=400
                    )

                det.cantidad = nueva_cantidad
                det.precio_unitario = nuevo_precio
                det.subtotal = det.cantidad * det.precio_unitario
                det.save()

                # 3) Descontar el nuevo stock
                Stock.objects.filter(producto=det.producto).update(
                    cantidad_actual=F("cantidad_actual") - nueva_cantidad
                )

            # 4) Guardar cambios de servicios
            for det in detalles_servicios:
                cant_str = request.POST.get(f"serv_cant_{det.id}")
                serv_id = request.POST.get(f"serv_servicio_{det.id}")
                pers_id = request.POST.get(f"serv_personal_{det.id}")

                if serv_id:
                    det.servicio = Servicio.objects.get(pk=serv_id)
                if pers_id:
                    det.colaborador_servicio = Personal.objects.get(pk=pers_id)
                if cant_str:
                    det.cantidad = int(cant_str)

                det.precio_unitario = det.servicio.precio
                det.subtotal = det.cantidad * det.precio_unitario
                det.save()

        return JsonResponse({"ok": True})

    # GET — anotar cada detalle con su stock para usarlo directo en el template
    for det in detalles_productos:
        det.stock_disponible = stock_por_producto.get(det.id, 0)

    # Lista completa de productos con su stock (para el select de cambio de producto)
    todos_productos = Producto.objects.all()
    todos_stock = []
    for p in todos_productos:
        entradas = DetalleCompra.objects.filter(producto=p).aggregate(t=Sum("cantidad"))["t"] or 0
        salidas = (
            DetalleVenta.objects.filter(producto=p, venta__estado="activa")
            .exclude(venta=venta)
            .aggregate(t=Sum("cantidad"))["t"] or 0
        )
        devuelto = (
            DetalleDevolucion.objects.filter(
                detalle_venta__producto=p,
                detalle_venta__venta__estado="activa",
            ).aggregate(t=Sum("cantidad_devuelta"))["t"] or 0
        )
        # Si este producto ya está en la venta, sumar lo que tiene
        det_actual = detalles_productos.filter(producto=p).first()
        ya_tiene = det_actual.cantidad_disponible_devolver if det_actual else 0
        s = entradas - salidas + devuelto + ya_tiene
        todos_stock.append({"producto": p, "stock": s})

    ctx = {
        "venta": venta,
        "detalles_productos": detalles_productos,
        "detalles_servicios": detalles_servicios,
        "servicios": servicios,
        "personal": personal,
        "todos_stock": todos_stock,
    }

    if es_ajax(request):
        html = render_to_string("ventas/form_editar_venta.html", ctx, request=request)
        return JsonResponse({"success": True, "html": html})

    return render(request, "ventas/form_editar_venta.html", ctx)

def detalle_venta_json(request, pk):
    venta = get_object_or_404(
        Venta.objects.select_related("cliente").prefetch_related(
            "detalles__producto",
            "detalles__servicio",
            "detalles__colaborador_servicio",
        ),
        pk=pk,
    )

    detalles = []
    for d in venta.detalles.all():
        nombre_item = ""
        tipo = ""

        if d.producto:
            nombre_item = d.producto.nombre
            tipo = "Producto"
        elif d.servicio:
            nombre_item = d.servicio.nombre
            tipo = "Servicio"

        detalles.append(
            {
                "tipo": tipo,
                "nombre": nombre_item,
                "cantidad": d.cantidad,
                "precio_unitario": str(d.precio_unitario),
                "subtotal": str(d.subtotal),
                "colaborador": str(d.colaborador_servicio)
                if d.colaborador_servicio
                else "",
            }
        )

    data = {
        "id": venta.id,
        "codigo_venta": venta.codigo_venta,
        "cliente": str(venta.cliente),
        "fecha": venta.fecha.strftime("%d/%m/%Y %H:%M") if venta.fecha else "",
        "estado": venta.estado,
        "total": str(venta.total),
        "detalles": detalles,
    }

    return JsonResponse(data)


@transaction.atomic
def anular_venta(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.prefetch_related("detalles__producto", "detalles__devoluciones"),
        id=venta_id
    )

    if venta.estado == "anulada":
        from django.urls import reverse
        return redirect(reverse("ventas:lista") + "?estado=activa")

    if request.method == "POST":
        for detalle in venta.detalles.all():
            if detalle.producto:
                cantidad_restante = detalle.cantidad_disponible_devolver
                if cantidad_restante > 0:
                    Stock.objects.filter(producto=detalle.producto).update(
                        cantidad_actual=F("cantidad_actual") + cantidad_restante
                    )
        venta.estado = "anulada"
        venta.save()

    from django.urls import reverse
    return redirect(reverse("ventas:lista") + "?estado=activa")


def clean(self):
    if self.cantidad <= 0:
        raise ValidationError("La cantidad debe ser mayor a 0")

    if self.precio_unitario <= 0:
        raise ValidationError("El precio debe ser mayor a 0")


# def reporte_ventas(request):
#     fecha_inicio = request.GET.get("fecha_inicio", "").strip()
#     fecha_fin = request.GET.get("fecha_fin", "").strip()
#     fecha_inicio_comp = request.GET.get("fecha_inicio_comp", "").strip()
#     fecha_fin_comp = request.GET.get("fecha_fin_comp", "").strip()

#     incluir_total = request.GET.get("incluir_total") == "1"
#     incluir_tabla = request.GET.get("incluir_tabla") == "1"
#     incluir_grafica = request.GET.get("incluir_grafica") == "1"
#     comparativo = request.GET.get("comparativo") == "1"

#     ventas = (
#         Venta.objects
#         .select_related("cliente")
#         .prefetch_related("detalles__producto", "detalles__servicio")
#         .annotate(
#             total_orden=Coalesce(
#                 Sum("detalles__subtotal"),
#                 Value(0),
#                 output_field=DecimalField(max_digits=12, decimal_places=2)
#             )
#         )
#         .order_by("fecha")
#     )

#     ventas_comp = Venta.objects.none()

#     if fecha_inicio and fecha_fin:
#         ventas = ventas.filter(fecha__date__range=[fecha_inicio, fecha_fin])

#     if comparativo and fecha_inicio_comp and fecha_fin_comp:
#         ventas_comp = (
#             Venta.objects
#             .select_related("cliente")
#             .prefetch_related("detalles__producto", "detalles__servicio")
#             .annotate(
#                 total_orden=Coalesce(
#                     Sum("detalles__subtotal"),
#                     Value(0),
#                     output_field=DecimalField(max_digits=12, decimal_places=2)
#                 )
#             )
#             .filter(fecha__date__range=[fecha_inicio_comp, fecha_fin_comp])
#             .order_by("fecha")
#         )

#     total_principal = ventas.aggregate(
#         total=Coalesce(
#             Sum("detalles__subtotal"),
#             Value(0),
#             output_field=DecimalField(max_digits=12, decimal_places=2)
#         )
#     )["total"]

#     total_comparativo = 0
#     if comparativo and fecha_inicio_comp and fecha_fin_comp:
#         total_comparativo = ventas_comp.aggregate(
#             total=Coalesce(
#                 Sum("detalles__subtotal"),
#                 Value(0),
#                 output_field=DecimalField(max_digits=12, decimal_places=2)
#             )
#         )["total"]

#     grafica_principal_labels = []
#     grafica_principal_data = []

#     if incluir_grafica:
#         grafica_principal = (
#             ventas
#             .annotate(dia=TruncDate("fecha"))
#             .values("dia")
#             .annotate(total=Coalesce(
#                 Sum("detalles__subtotal"),
#                 Value(0),
#                 output_field=DecimalField(max_digits=12, decimal_places=2)
#             ))
#             .order_by("dia")
#         )

#         grafica_principal_labels = [
#             item["dia"].strftime("%d/%m/%Y") for item in grafica_principal if item["dia"]
#         ]
#         grafica_principal_data = [float(item["total"]) for item in grafica_principal]

#     grafica_comp_labels = []
#     grafica_comp_data = []

#     if incluir_grafica and comparativo and fecha_inicio_comp and fecha_fin_comp:
#         grafica_comp = (
#             ventas_comp
#             .annotate(dia=TruncDate("fecha"))
#             .values("dia")
#             .annotate(total=Coalesce(
#                 Sum("detalles__subtotal"),
#                 Value(0),
#                 output_field=DecimalField(max_digits=12, decimal_places=2)
#             ))
#             .order_by("dia")
#         )

#         grafica_comp_labels = [
#             item["dia"].strftime("%d/%m/%Y") for item in grafica_comp if item["dia"]
#         ]
#         grafica_comp_data = [float(item["total"]) for item in grafica_comp]

#     return render(
#         request,
#         "ventas/reporte_ventas.html",
#         {
#             "ventas": ventas,
#             "ventas_comp": ventas_comp,
#             "fecha_inicio": fecha_inicio,
#             "fecha_fin": fecha_fin,
#             "fecha_inicio_comp": fecha_inicio_comp,
#             "fecha_fin_comp": fecha_fin_comp,
#             "incluir_total": incluir_total,
#             "incluir_tabla": incluir_tabla,
#             "incluir_grafica": incluir_grafica,
#             "comparativo": comparativo,
#             "total_principal": total_principal,
#             "total_comparativo": total_comparativo,
#             "grafica_principal_labels": grafica_principal_labels,
#             "grafica_principal_data": grafica_principal_data,
#             "grafica_comp_labels": grafica_comp_labels,
#             "grafica_comp_data": grafica_comp_data,
#         },
#     )
COLUMNAS_REPORTE_VENTAS = {
    "codigo_venta": "Código venta",
    "cliente": "Cliente",
    "productos_servicios": "Productos / Servicios",
    "fecha": "Fecha",
    "total": "Total",
    "estado": "Estado",
}


def construir_queryset_reporte_ventas(request):
    fecha_inicio = request.GET.get("fecha_inicio", "").strip()
    fecha_fin = request.GET.get("fecha_fin", "").strip()
    fecha_inicio_comp = request.GET.get("fecha_inicio_comp", "").strip()
    fecha_fin_comp = request.GET.get("fecha_fin_comp", "").strip()

    incluir_total = request.GET.get("incluir_total") == "1"
    incluir_grafica = request.GET.get("incluir_grafica") == "1"
    comparativo = request.GET.get("comparativo") == "1"
    tipo_grafica = request.GET.get("tipo_grafica", "bar").strip()
    columnas = request.GET.getlist("columnas")

    ventas = (
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__producto", "detalles__servicio")
        .annotate(
            total_orden=Coalesce(
                Sum("detalles__subtotal"),
                Value(0),
                output_field=DecimalField(max_digits=12, decimal_places=2),
            )
        )
        .order_by("fecha")
    )

    if fecha_inicio and fecha_fin:
        ventas = ventas.filter(fecha__date__range=[fecha_inicio, fecha_fin])

    ventas_comp = Venta.objects.none()
    if comparativo and fecha_inicio_comp and fecha_fin_comp:
        ventas_comp = (
            Venta.objects.select_related("cliente")
            .prefetch_related("detalles__producto", "detalles__servicio")
            .annotate(
                total_orden=Coalesce(
                    Sum("detalles__subtotal"),
                    Value(0),
                    output_field=DecimalField(max_digits=12, decimal_places=2),
                )
            )
            .filter(fecha__date__range=[fecha_inicio_comp, fecha_fin_comp])
            .order_by("fecha")
        )

    total_principal = ventas.aggregate(
        total=Coalesce(
            Sum("detalles__subtotal"),
            Value(0),
            output_field=DecimalField(max_digits=12, decimal_places=2),
        )
    )["total"]

    total_comparativo = 0
    if comparativo and fecha_inicio_comp and fecha_fin_comp:
        total_comparativo = ventas_comp.aggregate(
            total=Coalesce(
                Sum("detalles__subtotal"),
                Value(0),
                output_field=DecimalField(max_digits=12, decimal_places=2),
            )
        )["total"]

    grafica_principal_labels = []
    grafica_principal_data = []
    if incluir_grafica:
        grafica_principal = (
            ventas.annotate(dia=TruncDate("fecha"))
            .values("dia")
            .annotate(
                total=Coalesce(
                    Sum("detalles__subtotal"),
                    Value(0),
                    output_field=DecimalField(max_digits=12, decimal_places=2),
                )
            )
            .order_by("dia")
        )
        grafica_principal_labels = [
            item["dia"].strftime("%d/%m/%Y")
            for item in grafica_principal
            if item["dia"]
        ]
        grafica_principal_data = [float(item["total"]) for item in grafica_principal]

    grafica_comp_labels = []
    grafica_comp_data = []
    if incluir_grafica and comparativo and fecha_inicio_comp and fecha_fin_comp:
        grafica_comp = (
            ventas_comp.annotate(dia=TruncDate("fecha"))
            .values("dia")
            .annotate(
                total=Coalesce(
                    Sum("detalles__subtotal"),
                    Value(0),
                    output_field=DecimalField(max_digits=12, decimal_places=2),
                )
            )
            .order_by("dia")
        )
        grafica_comp_labels = [
            item["dia"].strftime("%d/%m/%Y") for item in grafica_comp if item["dia"]
        ]
        grafica_comp_data = [float(item["total"]) for item in grafica_comp]

    return {
        "ventas": ventas,
        "ventas_comp": ventas_comp,
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "fecha_inicio_comp": fecha_inicio_comp,
        "fecha_fin_comp": fecha_fin_comp,
        "incluir_total": incluir_total,
        "incluir_grafica": incluir_grafica,
        "comparativo": comparativo,
        "tipo_grafica": tipo_grafica,
        "columnas": columnas,
        "total_principal": total_principal,
        "total_comparativo": total_comparativo,
        "grafica_principal_labels": grafica_principal_labels,
        "grafica_principal_data": grafica_principal_data,
        "grafica_comp_labels": grafica_comp_labels,
        "grafica_comp_data": grafica_comp_data,
    }



def obtener_valor_columna_venta(venta, columna):
    if columna == "codigo_venta":
        return venta.codigo_venta
    if columna == "cliente":
        return str(venta.cliente)
    if columna == "productos_servicios":
        items = []
        for d in venta.detalles.all():
            nombre = ""
            if d.producto:
                nombre = d.producto.nombre
            elif d.servicio:
                nombre = d.servicio.nombre
            if nombre:
                items.append(f"{nombre} x{d.cantidad}")
        return ", ".join(items) if items else "Sin detalles"
    if columna == "fecha":
        return venta.fecha.strftime("%d/%m/%Y %H:%M") if venta.fecha else ""
    if columna == "total":
        return f"{venta.total:.2f}"
    if columna == "estado":
        return venta.estado.title()
    return ""


def vista_previa_reporte_ventas(request):
    data = construir_queryset_reporte_ventas(request)
    html = render_to_string(
        "ventas/partials/reporte_ventas_preview.html",
        {
            **data,
            "columnas_map": COLUMNAS_REPORTE_VENTAS,
        },
        request=request,
    )

    return JsonResponse(
        {
            "html": html,
            "grafica_principal_labels": data["grafica_principal_labels"],
            "grafica_principal_data": data["grafica_principal_data"],
            "grafica_comp_labels": data["grafica_comp_labels"],
            "grafica_comp_data": data["grafica_comp_data"],
            "incluir_grafica": data["incluir_grafica"],
            "comparativo": data["comparativo"],
            "tipo_grafica": data["tipo_grafica"],
            "fecha_inicio": data["fecha_inicio"],
            "fecha_fin": data["fecha_fin"],
            "fecha_inicio_comp": data["fecha_inicio_comp"],
            "fecha_fin_comp": data["fecha_fin_comp"],
        }
    )


def exportar_reporte_ventas(request):
    data = construir_queryset_reporte_ventas(request)
    formato = request.GET.get("formato", "pdf").strip().lower()
    columnas = data["columnas"] or [
        "codigo_venta",
        "cliente",
        "fecha",
        "total",
        "estado",
    ]

    if formato == "excel":
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte Ventas"

        encabezados = [COLUMNAS_REPORTE_VENTAS[c] for c in columnas]
        ws.append(encabezados)

        for cell in ws[1]:
            cell.font = Font(bold=True)

        for venta in data["ventas"]:
            ws.append([obtener_valor_columna_venta(venta, c) for c in columnas])

        if data["incluir_total"]:
            ws.append([])
            ws.append(["Total rango principal", f"{data['total_principal']:.2f}"])
            if data["comparativo"]:
                ws.append(
                    ["Total rango comparativo", f"{data['total_comparativo']:.2f}"]
                )

        if data["incluir_grafica"] and data["grafica_principal_labels"]:
            ws_chart = wb.create_sheet("Gráfica")
            ws_chart.append(["Fecha", "Total"])
            for lbl, val in zip(data["grafica_principal_labels"], data["grafica_principal_data"]):
                ws_chart.append([lbl, val])

            tipo = data.get("tipo_grafica", "bar")
            chart = LineChart() if tipo == "line" else BarChart()
            chart.title = "Ventas — rango principal"
            chart.y_axis.title = "Total ($)"
            chart.x_axis.title = "Fecha"
            chart.style = 10

            data_ref = Reference(ws_chart, min_col=2, min_row=1, max_row=len(data["grafica_principal_labels"]) + 1)
            cats_ref = Reference(ws_chart, min_col=1, min_row=2, max_row=len(data["grafica_principal_labels"]) + 1)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            chart.shape = 4
            ws_chart.add_chart(chart, "D2")

            if data["comparativo"] and data["grafica_comp_labels"]:
                ws_chart2 = wb.create_sheet("Gráfica comparativo")
                ws_chart2.append(["Fecha", "Total"])
                for lbl, val in zip(data["grafica_comp_labels"], data["grafica_comp_data"]):
                    ws_chart2.append([lbl, val])

                chart2 = LineChart() if tipo == "line" else BarChart()
                chart2.title = "Ventas — rango comparativo"
                chart2.y_axis.title = "Total ($)"
                chart2.style = 10
                data_ref2 = Reference(ws_chart2, min_col=2, min_row=1, max_row=len(data["grafica_comp_labels"]) + 1)
                cats_ref2 = Reference(ws_chart2, min_col=1, min_row=2, max_row=len(data["grafica_comp_labels"]) + 1)
                chart2.add_data(data_ref2, titles_from_data=True)
                chart2.set_categories(cats_ref2)
                ws_chart2.add_chart(chart2, "D2")

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="reporte_ventas.xlsx"'
        return response

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Reporte de Ventas", styles["Title"]))
    elements.append(Spacer(1, 12))
    elements.append(
        Paragraph(
            f"Rango principal: {data['fecha_inicio']} a {data['fecha_fin']}",
            styles["Normal"],
        )
    )
    if data["comparativo"]:
        elements.append(
            Paragraph(
                f"Rango comparativo: {data['fecha_inicio_comp']} a {data['fecha_fin_comp']}",
                styles["Normal"],
            )
        )
    elements.append(Spacer(1, 12))

    tabla_data = [[COLUMNAS_REPORTE_VENTAS[c] for c in columnas]]
    for venta in data["ventas"]:
        tabla_data.append([obtener_valor_columna_venta(venta, c) for c in columnas])

    tabla = Table(tabla_data, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#8d604a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.whitesmoke, colors.lightgrey],
                ),
            ]
        )
    )
    elements.append(tabla)

    if data["incluir_total"]:
        elements.append(Spacer(1, 12))
        elements.append(
            Paragraph(
                f"Total rango principal: {data['total_principal']:.2f}",
                styles["Heading3"],
            )
        )
        if data["comparativo"]:
            elements.append(
                Paragraph(
                    f"Total rango comparativo: {data['total_comparativo']:.2f}",
                    styles["Heading3"],
                )
            )

    if data["incluir_grafica"] and data["grafica_principal_labels"]:
        elements.append(Spacer(1, 16))
        elements.append(Paragraph("Gráfica de ventas — rango principal", styles["Heading3"]))
        elements.append(Spacer(1, 6))

        fig, ax = plt.subplots(figsize=(9, 3.5))
        tipo = data.get("tipo_grafica", "bar")
        labels = data["grafica_principal_labels"]
        valores = data["grafica_principal_data"]

        if tipo == "line":
            ax.plot(labels, valores, marker="o", color="#8d604a", linewidth=1.8)
        else:
            ax.bar(labels, valores, color="#8d604a")

        ax.set_ylabel("Total ($)")
        ax.tick_params(axis="x", rotation=45, labelsize=7)
        ax.tick_params(axis="y", labelsize=8)
        plt.tight_layout()

        img_buf = BytesIO()
        fig.savefig(img_buf, format="png", dpi=120)
        plt.close(fig)
        img_buf.seek(0)
        elements.append(RLImage(img_buf, width=560, height=220))

        if data["comparativo"] and data["grafica_comp_labels"]:
            elements.append(Spacer(1, 16))
            elements.append(Paragraph("Gráfica de ventas — rango comparativo", styles["Heading3"]))
            elements.append(Spacer(1, 6))

            fig2, ax2 = plt.subplots(figsize=(9, 3.5))
            labels2 = data["grafica_comp_labels"]
            valores2 = data["grafica_comp_data"]

            if tipo == "line":
                ax2.plot(labels2, valores2, marker="o", color="#4a7c8d", linewidth=1.8)
            else:
                ax2.bar(labels2, valores2, color="#4a7c8d")

            ax2.set_ylabel("Total ($)")
            ax2.tick_params(axis="x", rotation=45, labelsize=7)
            ax2.tick_params(axis="y", labelsize=8)
            plt.tight_layout()

            img_buf2 = BytesIO()
            fig2.savefig(img_buf2, format="png", dpi=120)
            plt.close(fig2)
            img_buf2.seek(0)
            elements.append(RLImage(img_buf2, width=560, height=220))

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="reporte_ventas.pdf"'
    return response

# ============================================================
# DEVOLUCIÓN DE VENTAS
# ============================================================


def devolucion_venta_json(request, venta_id):
    """
    GET: Devuelve en JSON los ítems de la venta con cuánto se puede devolver.
    """
    venta = get_object_or_404(
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__producto", "detalles__servicio", "detalles__devoluciones"),
        id=venta_id
    )

    # Validaciones de negocio
    if venta.estado == "anulada":
        return JsonResponse({"error": "No se puede devolver una venta anulada."}, status=400)

    items = []
    for d in venta.detalles.all():
        ya_devuelto = d.cantidad_devuelta
        disponible = d.cantidad_disponible_devolver

        nombre = ""
        tipo = ""
        if d.producto:
            nombre = d.producto.nombre
            tipo = "producto"
        elif d.servicio:
            nombre = d.servicio.nombre
            tipo = "servicio"

        items.append({
            "detalle_id": d.id,
            "nombre": nombre,
            "tipo": tipo,
            "precio_unitario": float(d.precio_unitario),
            "cantidad_original": d.cantidad,
            "ya_devuelto": ya_devuelto,
            "disponible": disponible,
        })

    return JsonResponse({
        "venta_id": venta.id,
        "codigo_venta": venta.codigo_venta,
        "cliente": str(venta.cliente),
        "total_venta": float(venta.total),
        "total_ya_devuelto": float(venta.total_devuelto),
        "items": items,
    })


@transaction.atomic
def registrar_devolucion(request, venta_id):
    """
    POST: Registra la devolución con validaciones completas.
    """
    if request.method != "POST":
        return JsonResponse({"error": "Método no permitido."}, status=405)

    venta = get_object_or_404(
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__devoluciones"),
        id=venta_id
    )

    # ── Validación 1: venta activa ──────────────────────────
    if venta.estado == "anulada":
        return JsonResponse({"error": "No se puede devolver una venta anulada."}, status=400)

    # ── Leer body JSON ──────────────────────────────────────
    try:
        body = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({"error": "Datos inválidos."}, status=400)

    motivo = body.get("motivo", "").strip()
    items_devolver = body.get("items", [])  # [{detalle_id, cantidad}]

    # ── Validación 2: motivo obligatorio ────────────────────
    if not motivo:
        return JsonResponse({"error": "El motivo de devolución es obligatorio."}, status=400)

    # ── Validación 3: al menos un ítem ─────────────────────
    if not items_devolver:
        return JsonResponse({"error": "Selecciona al menos un ítem para devolver."}, status=400)

    # ── Validación 4: ítem por ítem ─────────────────────────
    detalles_map = {d.id: d for d in venta.detalles.all()}
    lineas_validadas = []
    total_devolucion = Decimal("0")

    for item in items_devolver:
        detalle_id = item.get("detalle_id")
        try:
            cantidad = int(item.get("cantidad", 0))
        except (TypeError, ValueError):
            return JsonResponse({"error": f"Cantidad inválida en ítem {detalle_id}."}, status=400)

        if cantidad <= 0:
            continue  # el usuario dejó 0, ignorar silenciosamente

        # Existe en esta venta
        if detalle_id not in detalles_map:
            return JsonResponse({"error": f"El ítem {detalle_id} no pertenece a esta venta."}, status=400)

        detalle = detalles_map[detalle_id]
        disponible = detalle.cantidad_disponible_devolver

        # No superar disponible
        if cantidad > disponible:
            nombre = detalle.producto.nombre if detalle.producto else (
                detalle.servicio.nombre if detalle.servicio else f"ítem #{detalle_id}"
            )
            return JsonResponse({
                "error": f"'{nombre}': se intenta devolver {cantidad} pero solo hay {disponible} disponibles para devolución."
            }, status=400)

        subtotal = detalle.precio_unitario * cantidad
        total_devolucion += subtotal
        lineas_validadas.append((detalle, cantidad, subtotal))

    if not lineas_validadas:
        return JsonResponse({"error": "Ingresa al menos una cantidad mayor a 0."}, status=400)

    # ── Crear devolución ────────────────────────────────────
    devolucion = DevolucionVenta.objects.create(
        venta=venta,
        motivo=motivo,
        total_devuelto=total_devolucion,
    )

    for detalle, cantidad, subtotal in lineas_validadas:
        DetalleDevolucion.objects.create(
            devolucion=devolucion,
            detalle_venta=detalle,
            cantidad_devuelta=cantidad,
            subtotal_devuelto=subtotal,
        )
        # Restaurar stock si es producto
        if detalle.producto:
            Stock.objects.filter(producto=detalle.producto).update(
                cantidad_actual=F("cantidad_actual") + cantidad
            )

    # ── Anulación automática si todos los ítems fueron devueltos ──
    venta_anulada = False
    todos_devueltos = all(
        d.cantidad_disponible_devolver == 0
        for d in venta.detalles.all()
    )
    if todos_devueltos:
        venta.estado = "anulada"
        venta.save(update_fields=["estado"])
        venta_anulada = True

    return JsonResponse({
        "ok": True,
        "codigo_devolucion": devolucion.codigo_devolucion,
        "total_devuelto": float(total_devolucion),
        "venta_anulada": venta_anulada,
        "mensaje": f"Devolución {devolucion.codigo_devolucion} registrada por ${total_devolucion:.2f}.",
    })