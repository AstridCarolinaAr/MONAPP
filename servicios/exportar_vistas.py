# Funciones de exportación para el módulo servicios
from django.shortcuts import get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from io import BytesIO
from .models import Servicio


@login_required
def exportar_servicios_excel(request):
    """Exporta todos los servicios a un archivo Excel"""
    servicios = Servicio.objects.all()
    
    # Crear un nuevo libro de trabajo
    wb = Workbook()
    ws = wb.active
    ws.title = "Servicios"
    
    # Definir estilos
    header_fill = PatternFill(start_color="3a2a24", end_color="3a2a24", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # Headers
    headers = ['Nombre', 'Descripción', 'Precio', 'Estado', 'Fecha de Creación']
    ws.append(headers)
    
    # Aplicar estilos al header
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = border
    
    # Agregar datos
    for servicio in servicios:
        ws.append([
            servicio.nombre,
            servicio.descripcion or '',
            f"${float(servicio.precio):.2f}",
            'Activo' if servicio.activo else 'Inactivo',
            servicio.fecha_creacion.strftime('%d/%m/%Y') if servicio.fecha_creacion else ''
        ])
    
    # Ajustar ancho de columnas
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 18
    
    # Aplicar bordes a las celdas de datos
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='center')
    
    # Crear la respuesta
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="servicios_{datetime.now().strftime("%d%m%Y_%H%M%S")}.xlsx"'
    wb.save(response)
    
    return response


@login_required
def exportar_servicios_pdf(request):
    """Exporta todos los servicios a un archivo PDF"""
    servicios = Servicio.objects.all()
    
    # Crear el PDF en memoria
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer)
    story = []
    
    # Estilos
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#3a2a24'),
        spaceAfter=20,
        alignment=1  # Centrado
    )
    
    # Título
    story.append(Paragraph('Servicios', title_style))
    story.append(Spacer(1, 0.3))
    
    # Preparar datos de la tabla
    table_data = [['Nombre', 'Descripción', 'Precio', 'Estado', 'Fecha']]
    
    for servicio in servicios:
        table_data.append([
            servicio.nombre,
            servicio.descripcion or '',
            f"${float(servicio.precio):.2f}",
            'Activo' if servicio.activo else 'Inactivo',
            servicio.fecha_creacion.strftime('%d/%m/%Y') if servicio.fecha_creacion else ''
        ])
    
    # Crear tabla
    table = Table(table_data, colWidths=[1.5, 2.0, 1.0, 1.0, 1.0])
    
    # Estilos para la tabla
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3a2a24')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f0eb')]),
    ]))
    
    story.append(table)
    
    # Generar PDF
    doc.build(story)
    buffer.seek(0)
    
    # Crear respuesta
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="servicios_{datetime.now().strftime("%d%m%Y_%H%M%S")}.pdf"'
    
    return response
